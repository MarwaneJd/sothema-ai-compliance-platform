import io

import structlog

from app.core.exceptions import DocumentProcessingError

logger = structlog.get_logger()


class DocumentProcessor:
    """Extracts plain text from various document formats."""

    SUPPORTED_TYPES = {"pdf", "docx", "xlsx", "pptx", "txt", "md"}

    def extract_text(self, content: bytes, file_type: str) -> str:
        file_type = file_type.lower().strip(".")
        if file_type not in self.SUPPORTED_TYPES:
            raise DocumentProcessingError(f"Unsupported file type: {file_type}")

        try:
            extractor = getattr(self, f"_extract_{file_type}")
            text = extractor(content)
        except DocumentProcessingError:
            raise
        except Exception as e:
            logger.error("Text extraction failed", file_type=file_type, error=str(e))
            raise DocumentProcessingError(
                f"Failed to extract text from {file_type}: {e}"
            ) from e

        text = text.strip()
        if not text:
            raise DocumentProcessingError("Extracted text is empty")

        logger.info("Text extracted", file_type=file_type, char_count=len(text))
        return text

    def _extract_pdf(self, content: bytes) -> str:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)

    def _extract_docx(self, content: bytes) -> str:
        from docx import Document

        doc = Document(io.BytesIO(content))

        parts: list[str] = []

        # Headers first — regulatory docs often carry the document title,
        # number, and version in the header. Deduped (a linked header repeats
        # across sections).
        header_text = self._docx_header_footer_text(doc, "header")
        if header_text:
            parts.append(header_text)

        # Body in document order: paragraphs AND tables interleaved. python-docx's
        # `doc.paragraphs` skips table cells entirely, so the previous version
        # silently dropped all tabular content (acceptance criteria, limits,
        # responsibilities — exactly what SOPs put in tables). Walking the body
        # XML in order keeps a heading paragraph adjacent to the table beneath it,
        # so the RegulatoryChunker's section detection still works.
        for block in self._iter_docx_blocks(doc):
            block = block.strip()
            if block:
                parts.append(block)

        # Footers last — version / approval / page references.
        footer_text = self._docx_header_footer_text(doc, "footer")
        if footer_text:
            parts.append(footer_text)

        return "\n\n".join(parts)

    def _iter_docx_blocks(self, doc):
        """Yield body block text in document order — paragraph text and tables
        rendered as pipe-joined rows. Tables are the content `doc.paragraphs`
        misses."""
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, doc).text
            elif isinstance(child, CT_Tbl):
                yield self._docx_table_to_text(Table(child, doc))

    @staticmethod
    def _docx_table_to_text(table) -> str:
        """Render a table as one row per line, cells joined by ' | '. Collapses
        the repeated-cell duplication python-docx emits for merged cells."""
        lines: list[str] = []
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = cell.text.strip().replace("\n", " ")
                # Merged cells surface the same value on consecutive positions.
                if value and (not cells or cells[-1] != value):
                    cells.append(value)
            if cells:
                lines.append(" | ".join(cells))
        return "\n".join(lines)

    def _docx_header_footer_text(self, doc, which: str) -> str:
        """Collect unique header/footer paragraph + table text across all
        sections (linked sections repeat the same content)."""
        seen: set[str] = set()
        lines: list[str] = []
        for section in doc.sections:
            container = getattr(section, which, None)
            if container is None:
                continue
            for paragraph in container.paragraphs:
                text = paragraph.text.strip()
                if text and text not in seen:
                    seen.add(text)
                    lines.append(text)
            for table in container.tables:
                text = self._docx_table_to_text(table).strip()
                if text and text not in seen:
                    seen.add(text)
                    lines.append(text)
        return "\n".join(lines)

    def _extract_xlsx(self, content: bytes) -> str:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                sheets.append(f"[{sheet.title}]\n" + "\n".join(rows))
        wb.close()
        return "\n\n".join(sheets)

    def _extract_pptx(self, content: bytes) -> str:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(content))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            texts.append(text)
            if texts:
                slides.append(f"[Slide {i}]\n" + "\n".join(texts))
        return "\n\n".join(slides)

    def _extract_txt(self, content: bytes) -> str:
        return content.decode("utf-8", errors="replace")

    def _extract_md(self, content: bytes) -> str:
        return content.decode("utf-8", errors="replace")
